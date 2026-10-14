"""Build the sample sales data used to test Excel Bridge.

Writes two files from the same seeded data:
  ../../Sample Sales.xlsx     the workbook Adam opens in Excel
  ../dev/sample-data.json     the same rows for the browser preview (dev/dev.html)

Run:  python3 tools/make_sample.py   (from the Add-in folder)
"""

import datetime as dt
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

HERE = Path(__file__).resolve().parent
ADDIN = HERE.parent
WORKBOOK = ADDIN.parent / "Sample Sales.xlsx"
JSON_OUT = ADDIN / "dev" / "sample-data.json"

random.seed(20261014)

REGIONS = {
    "Europe": {"United Kingdom": 0.40, "Germany": 0.28, "France": 0.20, "Spain": 0.12},
    "North America": {"United States": 0.78, "Canada": 0.22},
    "Asia Pacific": {"Australia": 0.45, "Japan": 0.35, "Singapore": 0.20},
}
# Asia Pacific grows its share each year.
REGION_SHARE = {
    2023: {"Europe": 0.42, "North America": 0.38, "Asia Pacific": 0.20},
    2024: {"Europe": 0.40, "North America": 0.35, "Asia Pacific": 0.25},
    2025: {"Europe": 0.37, "North America": 0.31, "Asia Pacific": 0.32},
}
PRODUCTS = {
    "Furniture": {"Desk": 240, "Chair": 150, "Bookcase": 180, "Lamp": 45},
    "Technology": {"Laptop": 950, "Monitor": 260, "Headphones": 120, "Keyboard": 60},
    "Office Supplies": {"Notebook": 6, "Pens": 4, "Paper": 9, "Stapler": 12},
}
CATEGORY_SHARE = {"Furniture": 0.30, "Technology": 0.35, "Office Supplies": 0.35}
UNITS = {"Furniture": (1, 6), "Technology": (1, 8), "Office Supplies": (5, 60)}
MARGIN = {"Furniture": 0.14, "Technology": 0.19, "Office Supplies": 0.34}
# Online grows, Retail shrinks.
CHANNEL_SHARE = {
    2023: {"Online": 0.34, "Retail": 0.44, "Partner": 0.22},
    2024: {"Online": 0.42, "Retail": 0.37, "Partner": 0.21},
    2025: {"Online": 0.50, "Retail": 0.29, "Partner": 0.21},
}
ORDERS_PER_YEAR = {2023: 900, 2024: 1000, 2025: 1100}
SEASON = [0.80, 0.75, 0.90, 0.90, 0.95, 0.95, 0.90, 0.90, 1.00, 1.10, 1.30, 1.50]

HEADERS = ["Order ID", "Order Date", "Region", "Country", "Category", "Product",
           "Channel", "Units", "Revenue", "Profit"]
FORMATS = ["General", "dd/mm/yyyy", "General", "General", "General", "General",
           "General", "#,##0", "£#,##0.00", "£#,##0.00"]


def pick(weights):
    return random.choices(list(weights), weights=list(weights.values()))[0]


def make_rows():
    rows = []
    order_no = 10001
    for year, n in ORDERS_PER_YEAR.items():
        days = []
        for month in range(1, 13):
            first = dt.date(year, month, 1)
            last = (dt.date(year + (month == 12), month % 12 + 1, 1) - dt.timedelta(days=1)).day
            days += [(first.replace(day=d), SEASON[month - 1]) for d in range(1, last + 1)]
        dates = sorted(random.choices([d for d, _ in days], weights=[w for _, w in days], k=n))
        for date in dates:
            region = pick(REGION_SHARE[year])
            country = pick(REGIONS[region])
            category = pick(CATEGORY_SHARE)
            product = random.choice(list(PRODUCTS[category]))
            channel = pick(CHANNEL_SHARE[year])
            units = random.randint(*UNITS[category])
            price = PRODUCTS[category][product]
            discount = random.choice([0, 0, 0.05, 0.1, 0.2 if channel == "Partner" else 0.1])
            gross = units * price
            revenue = round(gross * (1 - discount), 2)
            profit = round(revenue * MARGIN[category] - gross * discount * 0.6
                           + random.uniform(-0.03, 0.03) * revenue, 2)
            rows.append([f"SO-{order_no}", date, region, country, category, product,
                         channel, units, revenue, profit])
            order_no += 1
    return rows


def excel_serial(date):
    return (date - dt.date(1899, 12, 30)).days


def write_json(rows):
    data = {
        "tableName": "Sales",
        "headers": HEADERS,
        "formats": FORMATS,
        "rows": [[r[0], excel_serial(r[1]), *r[2:]] for r in rows],
    }
    JSON_OUT.write_text(json.dumps(data, separators=(",", ":")))


def write_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    arial = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)

    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    last = len(rows) + 1
    total_row = last + 1
    for col, fmt in enumerate(FORMATS, start=1):
        for row in range(1, total_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = bold if row in (1, total_row) else arial
            if row > 1 and fmt != "General":
                cell.number_format = fmt

    ws.cell(row=total_row, column=1, value="Total")
    for col in (8, 9, 10):
        letter = "HIJ"[col - 8]
        ws.cell(row=total_row, column=col, value=f"=SUBTOTAL(109,{letter}2:{letter}{last})")

    table = Table(displayName="Sales", ref=f"A1:J{total_row}", totalsRowCount=1)
    columns = []
    for i, name in enumerate(HEADERS, start=1):
        column = TableColumn(id=i, name=name)
        if i == 1:
            column.totalsRowLabel = "Total"
        if i >= 8:
            column.totalsRowFunction = "sum"
        columns.append(column)
    table.tableColumns = columns
    table.autoFilter = AutoFilter(ref=f"A1:J{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    widths = [11, 12, 15, 16, 16, 12, 10, 8, 12, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEFGHIJ"[i]].width = w
    ws.freeze_panes = "A2"
    ws["A1"].comment = Comment("Sample data made up for testing Excel Bridge. Not real figures.",
                               "Excel Bridge")

    # Summary sheet: SUBTOTAL ignores rows hidden by a filter, so these follow the add-in.
    sm = wb.create_sheet("Summary")
    sm["A1"] = "Revenue by region, visible rows only"
    sm["A1"].font = Font(name="Arial", size=12, bold=True)
    sm["A2"] = ("These figures count only the rows the Sales table is showing. Turn on "
                "\"Filter the sheet too\" in Excel Bridge and they follow your clicks.")
    sm["A2"].font = arial
    sm["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sm.merge_cells("A2:D2")
    sm.row_dimensions[2].height = 42
    for col, label in enumerate(["Region", "Revenue", "Orders"], start=1):
        c = sm.cell(row=4, column=col, value=label)
        c.font = bold
    for i, region in enumerate(REGIONS, start=5):
        sm.cell(row=i, column=1, value=region).font = arial
        rev = (f"=SUMPRODUCT(SUBTOTAL(109,OFFSET(Sales!$I$2,ROW(Sales!$I$2:$I${last})"
               f"-ROW(Sales!$I$2),0)),--(Sales!$C$2:$C${last}=$A{i}))")
        cnt = (f"=SUMPRODUCT(SUBTOTAL(103,OFFSET(Sales!$C$2,ROW(Sales!$C$2:$C${last})"
               f"-ROW(Sales!$C$2),0)),--(Sales!$C$2:$C${last}=$A{i}))")
        sm.cell(row=i, column=2, value=rev)
        sm.cell(row=i, column=3, value=cnt)
    sm.cell(row=8, column=1, value="Total").font = bold
    sm.cell(row=8, column=2, value="=SUM(B5:B7)")
    sm.cell(row=8, column=3, value="=SUM(C5:C7)")
    for row in range(5, 9):
        sm.cell(row=row, column=2).number_format = "£#,##0"
        sm.cell(row=row, column=3).number_format = "#,##0"
        for col in (2, 3):
            sm.cell(row=row, column=col).font = bold if row == 8 else arial
    sm.column_dimensions["A"].width = 18
    sm.column_dimensions["B"].width = 14
    sm.column_dimensions["C"].width = 10
    sm.column_dimensions["D"].width = 30

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Revenue by region (visible rows)"
    chart.y_axis.title = None
    chart.legend = None
    chart.add_data(Reference(sm, min_col=2, min_row=4, max_row=7), titles_from_data=True)
    chart.set_categories(Reference(sm, min_col=1, min_row=5, max_row=7))
    chart.height = 7
    chart.width = 14
    sm.add_chart(chart, "A10")

    wb.calculation.fullCalcOnLoad = True
    wb.save(WORKBOOK)


if __name__ == "__main__":
    data = make_rows()
    write_json(data)
    write_workbook(data)
    print(f"{len(data)} rows -> {WORKBOOK.name}, {JSON_OUT.name}")
